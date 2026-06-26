import os
import time
import openpyxl
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any

class ReportingEngine:
    """
    Parses completed_clients.xlsx and failed_clients.xlsx outputs from amend_req.py,
    reconciles them with the in-memory runtime batch details, and generates a unified,
    beautifully-styled Final_Registration_Report.xlsx workbook.
    """
    @staticmethod
    def generate_final_report(
        batch_records: List[Dict[str, Any]],
        batch_id: str,
        completed_excel_path: Path,
        failed_excel_path: Path,
        output_report_path: Path,
        execution_durations: Dict[str, float] # Maps GSTIN -> duration in seconds
    ) -> None:
        """
        Synthesizes individual Excel result sheets into Final_Registration_Report.xlsx.
        """
        # Load successfully processed clients
        success_map = {}
        if completed_excel_path.exists():
            try:
                wb = openpyxl.load_workbook(completed_excel_path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 6:
                        gstin = str(row[0] or "").strip().upper()
                        arn = str(row[5] or "").strip()
                        notes = str(row[6] or "").strip()
                        timestamp = str(row[4] or "").strip()
                        success_map[gstin] = {
                            "arn": arn,
                            "remarks": notes or "Successfully processed bulk amendment.",
                            "timestamp": timestamp
                        }
                wb.close()
            except Exception:
                pass

        # Load failed processed clients
        failed_map = {}
        if failed_excel_path.exists():
            try:
                wb = openpyxl.load_workbook(failed_excel_path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 4:
                        gstin = str(row[0] or "").strip().upper()
                        reason = str(row[1] or "").strip()
                        ss_path = str(row[2] or "").strip()
                        timestamp = str(row[3] or "").strip()
                        failed_map[gstin] = {
                            "error": reason,
                            "screenshot": ss_path,
                            "timestamp": timestamp
                        }
                wb.close()
            except Exception:
                pass

        # Create fresh final report workbook
        wb_report = openpyxl.Workbook()
        ws_report = wb_report.active
        ws_report.title = "GST_Automation_Report"
        ws_report.sheet_view.showGridLines = True

        # Columns
        columns = [
            "GST_Reg_ID",
            "GSTIN",
            "Client",
            "Batch_ID",
            "Status",
            "Submission Time",
            "ARN Number",
            "Error",
            "Remarks",
            "Screenshot Path",
            "Processing Duration"
        ]

        # Style headers
        hdr_fill = openpyxl.styles.PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        hdr_font = openpyxl.styles.Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        hdr_align = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_name in enumerate(columns, 1):
            cell = ws_report.cell(row=1, column=col_idx, value=col_name)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = hdr_align
        ws_report.row_dimensions[1].height = 28

        # Style cells
        cell_font = openpyxl.styles.Font(name="Segoe UI", size=10)
        border_side = openpyxl.styles.Side(border_style="thin", color="D3D3D3")
        cell_border = openpyxl.styles.Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        row_idx = 2
        for record in batch_records:
            gst_reg_id = record.get("GST_Reg_ID") or record.get("GST_Reg_ID_hyperlink") or "N/A"
            # Strip standard format if required
            if isinstance(gst_reg_id, str) and " - " in gst_reg_id:
                gst_reg_id = gst_reg_id.split(" - ")[0]

            gstin = str(record.get("GSTIN") or "").strip().upper()
            client_name = record.get("Legal Name (GST Applicant)") or record.get("TradeName") or "Client"
            
            # Determine status & fetch parsed fields
            status = "PENDING"
            sub_time = "N/A"
            arn = "N/A"
            error_msg = "N/A"
            remarks = "Not processed during this run."
            screenshot = "N/A"
            
            duration_sec = execution_durations.get(gstin, 0.0)
            duration_str = f"{duration_sec:.1f}s" if duration_sec > 0.0 else "N/A"

            if gstin in success_map:
                status = "SUCCESS"
                sub_time = success_map[gstin]["timestamp"] or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                arn = success_map[gstin]["arn"]
                remarks = success_map[gstin]["remarks"]
            elif gstin in failed_map:
                status = "FAILED"
                sub_time = failed_map[gstin]["timestamp"] or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                error_msg = failed_map[gstin]["error"]
                screenshot = failed_map[gstin]["screenshot"]
                remarks = "Automation run failed."

            row_values = [
                gst_reg_id,
                gstin,
                client_name,
                batch_id,
                status,
                sub_time,
                arn,
                error_msg,
                remarks,
                screenshot,
                duration_str
            ]

            status_colors = {
                "SUCCESS": "D4EDDA", # Soft green
                "FAILED": "F8D7DA",  # Soft red
                "PENDING": "FFF3CD"  # Soft yellow
            }
            status_fill = openpyxl.styles.PatternFill(
                start_color=status_colors.get(status, "FFFFFF"),
                end_color=status_colors.get(status, "FFFFFF"),
                fill_type="solid"
            )

            for col_idx, val in enumerate(row_values, 1):
                cell = ws_report.cell(row=row_idx, column=col_idx, value=val)
                cell.font = cell_font
                cell.border = cell_border
                
                # Alignments
                if col_idx in [1, 2, 4, 5, 6, 7, 11]:
                    cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = openpyxl.styles.Alignment(horizontal="left", vertical="center")

                # Highlight Status column
                if col_idx == 5:
                    cell.fill = status_fill
                    cell.font = openpyxl.styles.Font(name="Segoe UI", size=10, bold=True)

                # Format Screenshot Path as a clickable link if it is an actual path
                if col_idx == 10 and val != "N/A" and val:
                    # Resolve relative path to absolute to prevent dead links
                    abs_path = os.path.abspath(val)
                    cell.value = "View Screenshot"
                    cell.hyperlink = abs_path
                    cell.font = openpyxl.styles.Font(name="Segoe UI", size=10, color="0000FF", underline="single")

            ws_report.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Adjust column widths dynamically
        for col in ws_report.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_report.column_dimensions[col_letter].width = max(14, min(45, max_len + 3))

        ws_report.freeze_panes = "A2"
        wb_report.save(output_report_path)
        wb_report.close()
